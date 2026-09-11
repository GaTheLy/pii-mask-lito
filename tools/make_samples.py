"""Create deterministic, wholly fictional sample files.

The samples exercise text, tabular, office-document, and raster inputs without
copying the content or layout of any source document. Names, identifiers, and
contact details are invented for this repository. The ``example.net`` domain
and ``555-01xx`` telephone range are reserved for examples.

Run: ``python -m tools.make_samples``
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "samples"

CONTACTS = (
    {
        "customer_id": "CUS-4827-9136",
        "name": "Mira Calder",
        "email": "mira.calder@example.net",
        "phone": "+1 202-555-0147",
        "street": "1847 Willowglass Lane",
        "city": "Port Haven, OR 97035",
        "plan": "Studio",
        "monthly_fee": 48.00,
        "credit": -8.00,
    },
    {
        "customer_id": "CUS-7314-2058",
        "name": "Dorian Vale",
        "email": "dorian.vale@example.net",
        "phone": "+1 202-555-0183",
        "street": "62 Juniper Arcade",
        "city": "North Ember, WA 98118",
        "plan": "Team",
        "monthly_fee": 125.00,
        "credit": -15.00,
    },
)


def note() -> None:
    person = CONTACTS[0]
    (OUT / "note.txt").write_text(
        "Customer support follow-up\n\n"
        f"Case ID: CASE-2048-7713\n"
        f"Customer: {person['name']}\n"
        f"Customer ID: {person['customer_id']}\n"
        f"Email: {person['email']}\n"
        f"Phone: {person['phone']}\n"
        f"Mailing address: {person['street']}, {person['city']}\n"
        "The customer asked to move the delivery window to 2026-10-03.\n",
        encoding="utf-8",
    )


def accounts_csv() -> None:
    rows = ["customer_id,name,email,phone,plan,monthly_fee,credit"]
    for person in CONTACTS:
        rows.append(
            f"{person['customer_id']},\"{person['name']}\",{person['email']},"
            f"\"{person['phone']}\",{person['plan']},{person['monthly_fee']:.2f},"
            f"{person['credit']:.2f}"
        )
    (OUT / "accounts.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")


def accounts_xlsx() -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Subscriptions"
    sheet.sheet_view.showGridLines = False
    sheet.append(
        [
            "Customer ID",
            "Name",
            "Email",
            "Phone",
            "Plan",
            "Monthly fee",
            "Credit",
            "Amount due",
        ]
    )
    for person in CONTACTS:
        sheet.append(
            [
                person["customer_id"],
                person["name"],
                person["email"],
                person["phone"],
                person["plan"],
                person["monthly_fee"],
                person["credit"],
                None,
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(2, len(CONTACTS) + 2):
        sheet[f"H{row}"] = f"=F{row}+G{row}"
        for column in range(1, 9):
            sheet.cell(row, column).font = Font(name="Arial", size=10)
            sheet.cell(row, column).alignment = Alignment(vertical="center")
        for column in (6, 7, 8):
            sheet.cell(row, column).number_format = '$#,##0.00'
    widths = (18, 18, 29, 18, 13, 15, 13, 15)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{len(CONTACTS) + 1}"
    sheet.print_area = f"A1:H{len(CONTACTS) + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.calculation.fullCalcOnLoad = True
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.created = datetime(2026, 1, 1)
    workbook.properties.modified = datetime(2026, 1, 1)
    workbook.save(OUT / "accounts.xlsx")


def review_docx() -> None:
    import docx
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    person = CONTACTS[1]
    document = docx.Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = title_style.element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    heading_style = document.styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style.font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph("Employee equipment handover", style="Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.line_spacing = 1.0
    document.add_paragraph(
        "This record lists the equipment assigned to a new team member and the "
        "contact details needed for delivery."
    )

    details = document.add_table(rows=0, cols=2)
    details.autofit = False
    details.columns[0].width = Inches(1.65)
    details.columns[1].width = Inches(4.9)
    fields = (
        ("Employee", person["name"]),
        ("Employee ID", "EMP-7305-A19"),
        ("Start date", "2026-10-12"),
        ("Email", person["email"]),
        ("Phone", person["phone"]),
        ("Delivery address", f"{person['street']}, {person['city']}"),
    )
    for index, (label, value) in enumerate(fields):
        cells = details.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F3F6F8" if index % 2 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(4)
        cells[0].paragraphs[0].runs[0].bold = True

    document.add_paragraph("Assigned items", style="Heading 1")
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(2.0)
    for cell, value in zip(table.rows[0].cells, ("Item", "Asset ID", "Status")):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "1F4E78")
        cell._tc.get_or_add_tcPr().append(shade)
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for values in (
        ("Laptop", "ASSET-4812", "Prepared"),
        ("Security key", "KEY-9307", "Prepared"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph("Recipient signature: ______________________________")

    for table_object in (details, table):
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "D9D9D9")
            borders.append(border)
        table_object._tbl.tblPr.append(borders)

    document.core_properties.title = "Employee equipment handover"
    document.core_properties.creator = ""
    document.core_properties.last_modified_by = ""
    document.save(OUT / "review.docx")


def screenshot() -> None:
    """Create a deterministic order-management panel as a raster fixture."""
    from PIL import Image, ImageDraw

    from pii_mask_lito.images import _font

    person = CONTACTS[0]
    image = Image.new("RGB", (1200, 560), "white")
    draw = ImageDraw.Draw(image)
    body, heading, small = _font(22), _font(30), _font(18)
    draw.rectangle([0, 0, 1200, 76], fill=(31, 78, 120))
    draw.text((34, 22), "Order details", fill="white", font=heading)
    draw.text((34, 104), "ORDER-8451-2206", fill=(31, 78, 120), font=heading)
    draw.text((34, 154), f"Customer   {person['name']}", fill="black", font=body)
    draw.text((34, 194), f"Customer ID   {person['customer_id']}", fill="black", font=body)
    draw.text((34, 234), f"Email   {person['email']}", fill="black", font=body)
    draw.text((34, 274), f"Phone   {person['phone']}", fill="black", font=body)
    draw.text((34, 314), f"Ship to   {person['street']}", fill="black", font=body)
    draw.text((136, 354), person["city"], fill="black", font=body)
    draw.rounded_rectangle([760, 130, 1152, 418], radius=12, fill=(243, 246, 248))
    draw.text((790, 158), "Order summary", fill=(31, 78, 120), font=body)
    draw.text((790, 210), "Desk lamp", fill="black", font=small)
    draw.text((1060, 210), "$64.00", fill="black", font=small)
    draw.text((790, 250), "Shipping", fill="black", font=small)
    draw.text((1068, 250), "$8.00", fill="black", font=small)
    draw.line([790, 296, 1120, 296], fill=(180, 188, 196), width=2)
    draw.text((790, 322), "Total", fill="black", font=body)
    draw.text((1048, 322), "$72.00", fill="black", font=body)
    draw.text((34, 474), "All details in this repository sample are fictional.", fill=(80, 88, 96), font=small)
    image.save(OUT / "screenshot.png", optimize=True)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    note()
    accounts_csv()
    accounts_xlsx()
    review_docx()
    screenshot()
    print(f"wrote {len(list(OUT.iterdir()))} synthetic samples to {OUT}")


if __name__ == "__main__":
    main()
