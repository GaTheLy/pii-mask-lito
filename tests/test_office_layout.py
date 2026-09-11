"""Layout invariants for the office maskers: `pytest tests/test_office_layout.py`.

A masked document that is unreadable, or that has lost its formulas, is a
failure as real as a leak -- nobody adopts a tool that returns a broken
spreadsheet. These checks build the smallest file that exhibits each structure,
mask it with a stub detector, and assert that everything except the PII came
back unchanged.

The detector is a stub on purpose: these tests are about where a tag lands and
what survives beside it, not about what the real detectors propose. Using the
real one would make the assertions depend on spaCy's mood.
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pii_mask_lito import office
from pii_mask_lito.model import Span
from pii_mask_lito.registry import TagRegistry

VALUES = {
    "Mira Calder": "PERSON",
    "Dorian Vale": "PERSON",
    "000-00-0000": "US_SSN",
    "mira.calder@example.net": "EMAIL_ADDRESS",
}


class StubDetector:
    """Flags a fixed list of literals wherever they appear. No model, no I/O."""

    def detect(self, tt, hint=""):
        spans = []
        for value, entity in VALUES.items():
            start = 0
            while True:
                at = tt.text.find(value, start)
                if at < 0:
                    break
                spans.append(Span(entity, at, at + len(value), 1.0, value))
                start = at + len(value)
        return sorted(spans, key=lambda s: s.start)


def mask(kind, build):
    """Build a file with `build`, mask it, and return (masked path, hits)."""
    tmp = Path(tempfile.mkdtemp(prefix="pii_mask_lito-layout-"))
    src, dest = str(tmp / f"in.{kind}"), str(tmp / f"out.{kind}")
    build(src)
    hits = getattr(office, f"mask_{kind}")(src, dest, StubDetector(), TagRegistry())
    return dest, hits


# --------------------------------------------------------------------------
# docx
# --------------------------------------------------------------------------


def test_docx_entity_split_across_runs_keeps_every_run_format():
    """Word can split "Mira Calder" across runs after an edit.

    The mask has to span all three, and the two runs that merely sit beside them
    must come back with their formatting intact. Rewriting the paragraph into
    run 0 masked the value but flattened bold, size and font across the whole
    paragraph -- a heading masked that way stops looking like a heading.
    """
    import docx
    from docx.shared import Pt

    def build(path):
        document = docx.Document()
        para = document.add_paragraph(style="Heading 1")
        para.add_run("Customer: ")
        for chunk in ("Mi", "ra Cal", "der"):
            para.add_run(chunk).bold = True
        tail = para.add_run(" seen today")
        tail.font.size = Pt(20)
        document.save(path)

    dest, hits = mask("docx", build)
    para = docx.Document(dest).paragraphs[0]
    assert [(s.text, tag) for s, tag in hits] == [("Mira Calder", "<NAME#0>")]
    assert para.text == "Customer: <NAME#0> seen today"
    assert para.style.name == "Heading 1", "paragraph style must survive"
    texts = [r.text for r in para.runs]
    assert texts[0] == "Customer: ", "the run before the entity is untouched"
    assert texts[1] == "<NAME#0>", "the tag inherits the run the entity started in"
    assert texts[-1] == " seen today", "the run after the entity is untouched"
    assert para.runs[1].bold is True, "the masked run keeps its own formatting"
    assert para.runs[-1].font.size == Pt(20), "sibling run formatting must survive"
    assert para.runs[0].bold is not True, "unrelated runs must not inherit bold"


def test_docx_hyperlink_text_is_masked_once():
    """Hyperlink text is part of paragraph.text but not of paragraph.runs.

    Masking into runs alone left the w:hyperlink untouched, so the paragraph
    printed the tag *and* the original: "Contact <NAME#0> todayDorian Vale".
    """
    import docx
    from docx.oxml.ns import nsdecls
    from docx.oxml.parser import parse_xml
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    def build(path):
        document = docx.Document()
        para = document.add_paragraph()
        para.add_run("Contact ")
        rid = document.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
        para._p.append(parse_xml(
            '<w:hyperlink %s r:id="%s"><w:r><w:t>Dorian Vale</w:t></w:r></w:hyperlink>'
            % (nsdecls("w", "r"), rid)
        ))
        para.add_run(" today")
        document.save(path)

    dest, _ = mask("docx", build)
    document = docx.Document(dest)
    assert document.paragraphs[0].text == "Contact <NAME#0> today"
    assert "Dorian Vale" not in office.text_of(dest)
    assert document.paragraphs[0]._p.xpath("w:hyperlink"), "the link itself must survive"


def test_docx_inline_image_survives_a_masked_paragraph():
    """A logo shares a paragraph with an account banner on letterhead.

    `run.text = ""` clears the run's children, so blanking runs to make room for
    a tag deleted the w:drawing with them and the letterhead came back blank.
    """
    import docx
    from PIL import Image

    def build(path):
        document = docx.Document()
        logo = str(Path(path).with_name("logo.png"))
        Image.new("RGB", (16, 16), (200, 0, 0)).save(logo)
        para = document.add_paragraph()
        para.add_run("Studio ")
        para.add_run().add_picture(logo)
        para.add_run(" Mira Calder")
        document.save(path)

    dest, _ = mask("docx", build)
    document = docx.Document(dest)
    assert document.paragraphs[0].text == "Studio  <NAME#0>"
    assert len(document.inline_shapes) == 1, "the inline image must survive masking"


def test_docx_masks_every_story_part():
    """Text box, first-page header, even-page header, footnote, and endnote."""
    import zipfile

    import docx
    from docx.oxml.ns import nsdecls
    from docx.oxml.parser import parse_xml

    def build(path):
        document = docx.Document()
        run = document.add_paragraph().add_run()
        run._r.append(parse_xml(
            '<w:pict %s><v:shape xmlns:v="urn:schemas-microsoft-com:vml" '
            'style="width:200pt;height:50pt"><v:textbox><w:txbxContent><w:p><w:r>'
            "<w:t>Box: Mira Calder</w:t></w:r></w:p></w:txbxContent></v:textbox>"
            "</v:shape></w:pict>" % nsdecls("w")
        ))
        section = document.sections[0]
        section.header.paragraphs[0].text = "Header Mira Calder"
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].text = "First page Dorian Vale"
        section.even_page_header.paragraphs[0].text = "Even page 000-00-0000"
        document.save(path)
        _add_notes(path)

    dest, _ = mask("docx", build)
    body = zipfile.ZipFile(dest)
    for name in body.namelist():
        if not name.endswith(".xml"):
            continue
        blob = body.read(name).decode("utf-8", "replace")
        for value in VALUES:
            assert value not in blob, f"{value} still readable in {name}"
    flat = office.text_of(dest)
    for label in ("Box:", "Header", "First page", "Even page", "Footnote:", "Endnote:"):
        assert label in flat, f"text_of must read the part holding {label!r}"


def test_docx_table_shape_and_nesting_survive():
    """Cell count, nesting depth and merges are layout, not content."""
    import docx

    def build(path):
        document = docx.Document()
        table = document.add_table(rows=3, cols=3)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Outer Mira Calder"
        table.cell(1, 1).add_table(rows=1, cols=2).cell(0, 0).text = "Nested 000-00-0000"
        table.cell(2, 0).merge(table.cell(2, 1))
        document.save(path)

    dest, hits = mask("docx", build)
    table = docx.Document(dest).tables[0]
    assert (len(table.rows), len(table.columns)) == (3, 3), "table dimensions must hold"
    assert table.style.name == "Table Grid"
    assert table.cell(0, 0).text == "Outer <NAME#0>"
    nested = table.cell(1, 1).tables
    assert len(nested) == 1 and len(nested[0].columns) == 2, "nested table must survive"
    assert nested[0].cell(0, 0).text == "Nested <SSN#0>"
    assert table.cell(2, 0)._tc is table.cell(2, 1)._tc, "the merge must still be one cell"
    # A merged cell is reachable from two row positions; its paragraph is still
    # one paragraph, so masking it twice would tag an already-tagged string.
    assert [tag for _, tag in hits].count("<NAME#0>") == 1


def test_docx_paragraph_text_matches_python_docx():
    """The offset walk must reproduce paragraph.text character for character.

    Detection runs on that string; if the walk counts a tab or a line break
    differently, every tag after it lands in the wrong place.
    """
    import docx

    document = docx.Document()
    para = document.add_paragraph()
    para.add_run("a\tb")
    para.add_run().add_break()
    para.add_run("c")
    assert office._paragraph_text(para._p) == para.text


# --------------------------------------------------------------------------
# xlsx
# --------------------------------------------------------------------------


def _add_notes(path):
    """Bolt a footnotes and an endnotes part onto a docx python-docx cannot make."""
    import shutil
    import zipfile

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    notes = {
        "word/footnotes.xml":
            f'<w:footnotes {w}><w:footnote w:id="1"><w:p><w:r>'
            "<w:t>Footnote: Dorian Vale</w:t></w:r></w:p></w:footnote></w:footnotes>",
        "word/endnotes.xml":
            f'<w:endnotes {w}><w:endnote w:id="1"><w:p><w:r>'
            "<w:t>Endnote: 000-00-0000</w:t></w:r></w:p></w:endnote></w:endnotes>",
    }
    types = (
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        '<Override PartName="/word/endnotes.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>'
    )
    rels = (
        '<Relationship Id="rIdFn9" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
        '<Relationship Id="rIdEn9" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/endnotes" Target="endnotes.xml"/>'
    )
    source = path + ".orig"
    shutil.move(path, source)
    with zipfile.ZipFile(source) as zin, \
            zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            blob = zin.read(name)
            # The package holds a JPEG thumbnail too: patch bytes, not text.
            if name == "[Content_Types].xml":
                blob = blob.replace(b"</Types>", types.encode() + b"</Types>")
            if name == "word/_rels/document.xml.rels":
                blob = blob.replace(b"</Relationships>", rels.encode() + b"</Relationships>")
            zout.writestr(name, blob)
        for name, blob in notes.items():
            zout.writestr(name, blob)


def test_xlsx_keeps_formulas_formats_widths_and_merges():
    """Everything that is not a string cell must come back byte-identical.

    The number format is what tells a reader 1234.5 is $1,234.50, the column
    width is what stops a masked cell rendering as ####, and a formula
    overwritten with a tag is a spreadsheet that no longer computes.
    """
    import openpyxl
    from openpyxl.styles import Font

    def build(path):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet["A1"], sheet["B1"], sheet["C1"] = "Name", "Amount", "Calc"
        sheet["A2"] = "Mira Calder"
        sheet["A2"].font = Font(bold=True, size=14)
        sheet["B2"] = 1234.5
        sheet["B2"].number_format = "$#,##0.00"
        sheet["C2"] = "=B2*2"
        sheet["C3"] = '=IF(A2="Mira Calder","dup","")'
        sheet["A4"] = "0055"
        sheet.column_dimensions["A"].width = 42.0
        sheet.row_dimensions[2].height = 33.0
        sheet.merge_cells("A6:C6")
        sheet["A6"] = "Banner Dorian Vale"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:C2"
        book.save(path)

    dest, _ = mask("xlsx", build)
    sheet = openpyxl.load_workbook(dest).active
    assert sheet["A2"].value == "<NAME#0>"
    assert sheet["B2"].value == 1234.5 and sheet["B2"].data_type == "n"
    assert sheet["B2"].number_format == "$#,##0.00", "number format must survive"
    assert sheet["C2"].value == "=B2*2" and sheet["C2"].data_type == "f"
    assert sheet["C3"].value == '=IF(A2="<NAME#0>","dup","")', "literal masked, syntax kept"
    assert sheet["C3"].data_type == "f", "a masked formula must still be a formula"
    assert sheet["A4"].value == "0055", "an untouched string keeps its leading zero"
    assert sheet["A2"].font.bold and sheet["A2"].font.size == 14.0
    assert sheet.column_dimensions["A"].width == 42.0
    assert sheet.row_dimensions[2].height == 33.0
    assert [str(r) for r in sheet.merged_cells.ranges] == ["A6:C6"]
    assert sheet["A6"].value == "Banner <NAME#1>", "the anchor of a merge is a normal cell"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:C2"


def test_xlsx_formula_references_survive_a_renamed_tab():
    """A tab named after a person is a disclosure; a #REF! is a broken file.

    Renaming the sheet without repointing its references gives both at once.
    """
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    def build(path):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Mira Calder projects"
        sheet["A2"] = 7
        summary = book.create_sheet("Summary")
        summary["A1"] = "='Mira Calder projects'!A2"
        summary["A2"] = "=SUM('Mira Calder projects'!A2:A2)+1"
        book.defined_names.add(DefinedName("MyRef", attr_text="'Mira Calder projects'!$A$2"))
        book.save(path)

    dest, _ = mask("xlsx", build)
    book = openpyxl.load_workbook(dest)
    assert book.sheetnames == ["<NAME#0> projects", "Summary"]
    summary = book["Summary"]
    assert summary["A1"].value == "='<NAME#0> projects'!A2"
    assert summary["A2"].value == "=SUM('<NAME#0> projects'!A2:A2)+1"
    assert book.defined_names["MyRef"].value == "'<NAME#0> projects'!$A$2"
    for row in summary.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                name = cell.value.split("'")[1]
                assert name in book.sheetnames, f"{name} is a #REF! waiting to happen"


def test_xlsx_masks_comments_links_and_print_headers():
    """Three surfaces that a cell-value sweep cannot see, all readable in Excel."""
    import openpyxl
    from openpyxl.comments import Comment

    def build(path):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet["A1"] = "note"
        sheet["A1"].comment = Comment("Confirm with Dorian Vale", "auditor")
        sheet["A2"] = "email"
        sheet["A2"].hyperlink = "mailto:mira.calder@example.net"
        sheet.oddHeader.center.text = "Chart for Mira Calder"
        sheet.oddFooter.right.text = "SSN 000-00-0000"
        book.save(path)

    dest, _ = mask("xlsx", build)
    sheet = openpyxl.load_workbook(dest).active
    assert sheet["A1"].comment.text == "Confirm with <NAME#0>"
    assert sheet["A1"].comment.author == "auditor", "comment metadata must survive"
    assert sheet["A2"].hyperlink.target == "mailto:<EMAIL#0>"
    assert sheet.oddHeader.center.text == "Chart for <NAME#1>"
    assert sheet.oddFooter.right.text == "SSN <SSN#0>"
    flat = office.text_of(dest)
    for value in VALUES:
        assert value not in flat, f"{value} still readable through text_of"


def test_xlsx_charts_and_images_survive_the_round_trip():
    """openpyxl rewrites the file from its own model; anything it drops is lost."""
    import zipfile

    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image

    def build(path):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet["A1"], sheet["B1"] = "Name", "Count"
        sheet["A2"], sheet["B2"] = "Mira Calder", 3
        chart = BarChart()
        data = Reference(sheet, min_col=2, min_row=1, max_row=2)
        chart.add_data(data, titles_from_data=True)
        sheet.add_chart(chart, "E5")
        logo = str(Path(path).with_name("logo.png"))
        Image.new("RGB", (16, 16), (0, 0, 200)).save(logo)
        sheet.add_image(XLImage(logo), "H5")
        book.save(path)
        return path

    dest, _ = mask("xlsx", build)
    names = zipfile.ZipFile(dest).namelist()
    assert any(n.startswith("xl/charts/") for n in names), "the chart must survive"
    assert any(n.startswith("xl/media/") for n in names), "the image must survive"


def test_document_properties_are_masked_in_both_formats():
    """dc:creator and dc:title ship with the file and no reader has to open it.

    Both can remain readable in docProps/core.xml after every paragraph and cell
    has been masked because export tools can populate the author and title.
    """
    import docx
    import openpyxl

    def build_docx(path):
        document = docx.Document()
        document.add_paragraph("body")
        document.core_properties.author = "Mira Calder"
        document.core_properties.title = "Project summary for Dorian Vale"
        document.save(path)

    def build_xlsx(path):
        book = openpyxl.Workbook()
        book.active["A1"] = "body"
        book.properties.creator = "Mira Calder"
        book.properties.title = "Project status for Dorian Vale"
        book.save(path)

    dest, _ = mask("docx", build_docx)
    properties = docx.Document(dest).core_properties
    assert properties.author == "<NAME#0>"
    assert properties.title == "Project summary for <NAME#1>"

    dest, _ = mask("xlsx", build_xlsx)
    properties = openpyxl.load_workbook(dest).properties
    assert properties.creator == "<NAME#0>"
    assert properties.title == "Project status for <NAME#1>"
    assert "Dorian Vale" not in office.text_of(dest), "verification must read properties"


def main():
    checks = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for check in checks:
        check()
        print(f"  ok  {check.__name__}")
    print(f"\n{len(checks)} checks passed")


if __name__ == "__main__":
    main()
