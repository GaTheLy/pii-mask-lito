"""Flat and structured text formats: txt, csv, docx, xlsx.

No geometry here -- tokens carry no bbox, so detection runs on text alone and
replacement is a string substitution that preserves the surrounding container
(the run, the cell, the row).

The container is the whole difficulty. An office file is not one text stream: a
docx keeps visible words in six different parts and inside text boxes, and an
xlsx keeps them in comments, tab names and print headers as well as cells. Every
surface this module does not walk is a surface that leaves the building
unmasked, so the walks below are deliberately exhaustive rather than convenient.
"""

from __future__ import annotations

import csv
import io
import re

from .model import TokenText


def _tagged(text: str, detector, registry, context: str = "") -> list:
    """Detect in a string and pair every span with the tag that replaces it.

    Split out of `replace` because the docx path needs the spans themselves --
    it splices tags into individual `w:t` nodes rather than into one string.
    """
    tt = TokenText.from_text(text)
    spans = detector.detect(tt, hint=context)
    return [(s, registry.tag(s.entity, s.text)) for s in sorted(spans, key=lambda s: s.start)]


def replace(text: str, detector, registry, context: str = "") -> tuple[str, list]:
    """Detect in a string and return it with every span swapped for its tag.

    `context` is a label hint passed alongside the text, never spliced into it.
    It carries a column header into a spreadsheet cell so an otherwise
    ambiguous value under an identifier label can be recognized. Passing it as
    a hint rather than a text prefix keeps offsets aligned with the cell value.
    """
    hits = _tagged(text, detector, registry, context)
    if not hits:
        return text, []
    out, cursor = [], 0
    for span, tag in hits:
        out.append(text[cursor : span.start])
        out.append(tag)
        cursor = span.end
    out.append(text[cursor:])
    return "".join(out), hits


def mask_txt(src: str, dest: str, detector, registry) -> list:
    text = open(src, encoding="utf-8", errors="replace").read()
    masked, applied = replace(text, detector, registry)
    open(dest, "w", encoding="utf-8").write(masked)
    return applied


def mask_csv(src: str, dest: str, detector, registry) -> list:
    applied = []
    with open(src, newline="", encoding="utf-8", errors="replace") as fh:
        rows = list(csv.reader(fh))
    header = rows[0] if rows else []
    for index, row in enumerate(rows):
        for i, cell in enumerate(row):
            # Row 0 is treated as labels, not data, and supplies context below.
            context = "" if index == 0 else (header[i] if i < len(header) else "")
            row[i], hits = replace(cell, detector, registry, context=context)
            applied += hits
    with open(dest, "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)
    return applied


# --------------------------------------------------------------------------
# docx
# --------------------------------------------------------------------------

# Every part of a .docx that can hold a paragraph a reader will see. Walking
# `document.paragraphs`, `document.tables` and `section.header/.footer` -- the
# obvious four -- misses all of these, each verified leaking on a built sample:
# a first-page or even-page header lives in its own part that `section.header`
# never returns, and footnotes, endnotes and comments are separate parts
# entirely. Matched on the content-type suffix so the check does not depend on
# where python-docx happens to file each part in its object model.
_STORY_TYPES = (
    "wordprocessingml.document.main+xml",
    "wordprocessingml.document.glossary+xml",
    "wordprocessingml.header+xml",
    "wordprocessingml.footer+xml",
    "wordprocessingml.footnotes+xml",
    "wordprocessingml.endnotes+xml",
    "wordprocessingml.comments+xml",
)

# The run-level children whose text python-docx folds into `paragraph.text`.
# Detection ran on that string, so the offset walk must reproduce it exactly or
# every tag lands a few characters off. Only `w:t` holds editable characters: a
# tab or a break contributes one character that has to be counted, but deleting
# it to make room for a tag would reflow the page.
_RUN_CONTENT = "w:br | w:cr | w:noBreakHyphen | w:ptab | w:t | w:tab"


def _pieces(paragraph) -> list[tuple]:
    """Text-carrying leaves of one paragraph, in reading order.

    Walks `w:hyperlink` as well as `w:r`, exactly as python-docx's own
    `paragraph.text` does. That is not a detail: the old code rewrote the whole
    masked paragraph into `runs[0]`, and `runs` excludes runs nested in a
    hyperlink, so a paragraph reading "Contact Dorian Vale today" could come
    out as "Contact <NAME#0> todayDorian Vale" -- the mask applied and the original still
    printed beside it.

    Text inside `w:txbxContent` is deliberately not collected: those paragraphs
    are visited in their own right by the `w:p` sweep in `mask_docx`, and a run
    holding a text box contributes no characters to the paragraph around it.
    """
    from docx.oxml.ns import qn

    out = []
    for child in paragraph.xpath("w:r | w:hyperlink"):
        runs = [child] if child.tag == qn("w:r") else child.xpath("w:r")
        for run in runs:
            for node in run.xpath(_RUN_CONTENT):
                out.append((node, str(node), node.tag == qn("w:t")))
    return out


def _paragraph_text(paragraph) -> str:
    return "".join(text for _, text, _ in _pieces(paragraph))


def _set_text(node, text: str) -> None:
    from docx.oxml.ns import qn

    node.text = text
    # Word collapses leading and trailing spaces in a `w:t` unless told not to,
    # so a label followed by a tag retains its intentional spacing.
    if text != text.strip():
        node.set(qn("xml:space"), "preserve")


def _mask_paragraph(paragraph, detector, registry) -> list:
    """Mask one paragraph in place, touching only the runs an entity covers.

    Each tag is emitted at the character where its span starts, so it inherits
    the formatting of the run the entity started in, and the rest of the covered
    characters are deleted wherever they live. An entity split across runs --
    "Jo" + "hn Sm" + "ith", which is what Word produces after a spellcheck --
    is therefore masked without disturbing any other run.

    Writing the whole paragraph into one run would discard formatting, embedded
    drawings, and hyperlink structure, so edits stay at the text-leaf level.
    """
    pieces = _pieces(paragraph)
    text = "".join(t for _, t, _ in pieces)
    if not text.strip():
        return []
    hits = _tagged(text, detector, registry)
    if not hits:
        return []
    tag_at = {span.start: tag for span, tag in hits}
    covered = set()
    for span, _ in hits:
        covered.update(range(span.start, span.end))
    cursor = 0
    pending: list[str] = []
    last_editable = None
    for node, piece, editable in pieces:
        start, cursor = cursor, cursor + len(piece)
        if not editable:
            # A span that starts on a tab or a line break has no `w:t` of its
            # own to hold the tag. Carry it to the next one; dropping it would
            # delete the value from the report's own output.
            pending += [tag_at[i] for i in range(start, cursor) if i in tag_at]
            continue
        out, pending = pending, []
        for i in range(start, cursor):
            if i in tag_at:
                out.append(tag_at[i])
            if i not in covered:
                out.append(text[i])
        _set_text(node, "".join(out))
        last_editable = node
    if pending and last_editable is not None:
        _set_text(last_editable, str(last_editable) + "".join(pending))
    return hits


# Document properties travel with the file and are the first thing an audit
# tool prints. Creator, title, subject, and comments may all contain personal
# data even after visible paragraphs and cells have been masked. The two
# libraries spell equivalent fields differently, hence two lists.
_DOCX_PROPERTIES = ("author", "last_modified_by", "title", "subject", "comments",
                    "keywords", "category")
_XLSX_PROPERTIES = ("creator", "lastModifiedBy", "title", "subject", "description",
                    "keywords", "category")


def _mask_properties(properties, names, detector, registry) -> list:
    applied = []
    for name in names:
        value = getattr(properties, name, None)
        # created/modified are datetimes, revision an int: only text is masked.
        if not isinstance(value, str) or not value.strip():
            continue
        masked, hits = replace(value, detector, registry)
        if hits:
            setattr(properties, name, masked)
            applied += hits
    return applied


def _property_text(properties, names) -> list[str]:
    return [
        value for value in (getattr(properties, name, None) for name in names)
        if isinstance(value, str) and value.strip()
    ]


def _story_parts(document):
    """Every part of the package that can hold a visible paragraph."""
    return [
        part
        for part in document.part.package.iter_parts()
        if any(part.content_type.endswith(suffix) for suffix in _STORY_TYPES)
    ]


def _story_root(part):
    """Root element of a story part, and whether it had to be parsed by hand.

    python-docx models the document, headers and footers, so those parts expose
    a live `element` that `document.save` reserializes. Footnotes, endnotes and
    comments are not in its object model on every version and load as opaque
    blobs; parsing and writing the blob back prevents hidden story text from
    bypassing masking.
    """
    from docx.oxml.parser import parse_xml

    root = getattr(part, "element", None)
    if root is not None:
        return root, False
    return parse_xml(part.blob), True


def mask_docx(src: str, dest: str, detector, registry) -> list:
    import docx
    from docx.opc.oxml import serialize_part_xml
    from docx.oxml.ns import qn

    document = docx.Document(src)
    applied = []
    for part in _story_parts(document):
        root, owned = _story_root(part)
        # Every `w:p` in the part, at any depth. That one sweep covers body
        # paragraphs, table cells nested to any depth, and -- the case the
        # old container walk could not reach at all -- paragraphs inside a
        # text box, whose content lives under a run rather than under the body.
        for paragraph in root.iter(qn("w:p")):
            applied += _mask_paragraph(paragraph, detector, registry)
        if owned:
            part._blob = serialize_part_xml(root)
    applied += _mask_properties(document.core_properties, _DOCX_PROPERTIES, detector, registry)
    document.save(dest)
    return applied


# --------------------------------------------------------------------------
# xlsx
# --------------------------------------------------------------------------

# Excel reads a leading "=" as syntax, not text, so only the double-quoted
# literals inside a formula may be rewritten. Masking the rest turns
# ='Calder Mira'!A2 into ='<NAME#0>'!A2 and every cell reading it shows #REF!.
# The doubled quote ("") is Excel's own escape and stays inside the literal.
_XL_LITERAL = re.compile(r'"(?:[^"]|"")*"')

# Excel's tab-name ceiling. A longer title is written but flagged unreadable by
# some applications, and a masked title is easily longer than the name it
# replaced.
_XL_TITLE_MAX = 31


def _mask_formula(formula: str, detector, registry, context: str = "") -> tuple[str, list]:
    out, applied, cursor = [], [], 0
    for m in _XL_LITERAL.finditer(formula):
        out.append(formula[cursor : m.start()])
        masked, hits = replace(m.group(0)[1:-1], detector, registry, context)
        out.append('"%s"' % masked)
        applied += hits
        cursor = m.end()
    out.append(formula[cursor:])
    return "".join(out), applied


def _mask_cell(cell, detector, registry, context: str) -> list:
    from openpyxl.comments import Comment

    applied = []
    value = cell.value
    if isinstance(value, str) and value.strip():
        if cell.data_type == "f":
            masked, hits = _mask_formula(value, detector, registry, context)
        else:
            masked, hits = replace(value, detector, registry, context)
        # Assigned only on a hit. Writing an unchanged string back re-types the
        # cell from openpyxl's inference rather than from the file it came out
        # of, which is a free way to lose a formula or a leading zero.
        if hits:
            cell.value = masked
            applied += hits
    # A cell comment is a sticky note Excel shows on hover. A person's name can
    # survive any sweep that only reads cell values because comments live in a
    # separate XML part.
    comment = getattr(cell, "comment", None)
    if comment is not None and comment.text:
        masked, hits = replace(comment.text, detector, registry, context)
        if hits:
            cell.comment = Comment(masked, comment.author, comment.height, comment.width)
            applied += hits
    # A hyperlink target is not shown in the grid but ships in the sheet rels:
    # mailto:mira.calder@example.net is a name and address in one string.
    link = getattr(cell, "hyperlink", None)
    if link is not None:
        for attr in ("target", "location", "display", "tooltip"):
            text = getattr(link, attr, None)
            if not isinstance(text, str) or not text.strip():
                continue
            masked, hits = replace(text, detector, registry, context)
            if hits:
                setattr(link, attr, masked)
                applied += hits
    return applied


def _mask_print_headers(sheet, detector, registry) -> list:
    """Mask the banner Excel stamps on every printed page.

    Print headers live in sheet XML rather than cells and can be the only place
    a personal value appears, so the grid sweep alone is insufficient.
    """
    applied = []
    blocks = (
        sheet.oddHeader, sheet.oddFooter,
        sheet.evenHeader, sheet.evenFooter,
        sheet.firstHeader, sheet.firstFooter,
    )
    for block in blocks:
        for side in (block.left, block.center, block.right):
            if not side.text:
                continue
            masked, hits = replace(side.text, detector, registry)
            if hits:
                side.text = masked
                applied += hits
    return applied


def _mask_titles(book, detector, registry) -> list:
    """Mask sheet tab names, then repoint every formula that named one.

    A tab reading "Calder, Mira - projects" is as much a disclosure as a cell, but
    renaming a sheet does not rewrite the formulas that reference it: openpyxl
    leaves ='Calder, Mira - projects'!A2 pointing at a sheet that no longer exists.
    So the references are rewritten here, in the same pass, and never by the
    formula masker -- which is exactly why that masker refuses to touch
    anything outside a quoted literal.
    """
    applied, renames = [], {}
    taken = {s.title.casefold() for s in book.worksheets}
    for sheet in book.worksheets:
        masked, hits = replace(sheet.title, detector, registry)
        if not hits or masked == sheet.title:
            continue
        masked = masked[:_XL_TITLE_MAX]
        # Two tabs can mask to one name despite whitespace differences and share
        # a registry index, and truncation collides too. openpyxl raises on a
        # duplicate title, which would abort the whole job over a tab name.
        if masked.casefold() in taken:
            stem, suffix = masked, 2
            while masked.casefold() in taken:
                masked = "%s %d" % (stem[: _XL_TITLE_MAX - 3], suffix)
                suffix += 1
        taken.add(masked.casefold())
        renames[sheet.title] = masked
        applied += hits
    if not renames:
        return applied
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                cell.value = _repoint(cell.value, renames)
    for name in book.defined_names.values():
        if isinstance(getattr(name, "value", None), str):
            name.value = _repoint(name.value, renames)
    for old, new in renames.items():
        book[old].title = new
    return applied


def _repoint(formula: str, renames: dict) -> str:
    """Swap old sheet names for new ones in a reference, quoted or bare.

    Only a name followed by "!" is touched, so a sheet called "Labs" cannot
    corrupt the word Labs inside a string literal or a function name.
    """
    for old, new in renames.items():
        formula = formula.replace("'%s'!" % old, "'%s'!" % new)
        if not re.search(r"[\s'\"]", old):
            formula = re.sub(r"(?<![A-Za-z0-9_.])%s!" % re.escape(old), "%s!" % new, formula)
    return formula


def mask_xlsx(src: str, dest: str, detector, registry) -> list:
    import openpyxl

    book = openpyxl.load_workbook(src)
    applied = []
    for sheet in book.worksheets:
        header: list[str] = []
        for index, row in enumerate(sheet.iter_rows()):
            if index == 0:
                header = [str(c.value or "") for c in row]
            for i, cell in enumerate(row):
                context = "" if index == 0 else (header[i] if i < len(header) else "")
                applied += _mask_cell(cell, detector, registry, context)
        applied += _mask_print_headers(sheet, detector, registry)
    applied += _mask_titles(book, detector, registry)
    applied += _mask_properties(book.properties, _XLSX_PROPERTIES, detector, registry)
    book.save(dest)
    return applied


def text_of(path: str) -> str:
    """Flat text of any office format, for read-back verification.

    Reads every surface the maskers write, and for the same reason: a
    verification pass that only reads cell values and body paragraphs cannot
    fail on a name left in a footnote, a text box, a tab name or a comment, so
    the gate could pass a document that still contains a personal name.
    """
    suffix = path.rsplit(".", 1)[-1].lower()
    if suffix == "docx":
        import docx
        from docx.oxml.ns import qn

        document = docx.Document(path)
        parts = []
        for part in _story_parts(document):
            root, _ = _story_root(part)
            parts += [_paragraph_text(p) for p in root.iter(qn("w:p"))]
        parts += _property_text(document.core_properties, _DOCX_PROPERTIES)
        return "\n".join(parts)
    if suffix == "xlsx":
        import openpyxl

        book = openpyxl.load_workbook(path)
        buf = io.StringIO()
        for sheet in book.worksheets:
            buf.write(sheet.title + "\n")
            for row in sheet.iter_rows():
                buf.write(" ".join(str(c.value) for c in row if c.value is not None) + "\n")
                for cell in row:
                    comment = getattr(cell, "comment", None)
                    if comment is not None and comment.text:
                        buf.write(comment.text + "\n")
                    link = getattr(cell, "hyperlink", None)
                    for attr in ("target", "location", "display", "tooltip"):
                        text = getattr(link, attr, None) if link is not None else None
                        if isinstance(text, str) and text.strip():
                            buf.write(text + "\n")
            for block in (sheet.oddHeader, sheet.oddFooter, sheet.evenHeader,
                          sheet.evenFooter, sheet.firstHeader, sheet.firstFooter):
                for side in (block.left, block.center, block.right):
                    if side.text:
                        buf.write(side.text + "\n")
        for value in _property_text(book.properties, _XLSX_PROPERTIES):
            buf.write(value + "\n")
        return buf.getvalue()
    return open(path, encoding="utf-8", errors="replace").read()
